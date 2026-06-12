from __future__ import annotations

import json
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]


def cell_value(v):
    if v is None:
        return ""
    return str(v).strip()


def inspect_xlsx(path: Path) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    sheets = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 12), values_only=True):
            rows.append([cell_value(v) for v in row[:12]])
        sheets.append(
            {
                "title": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "preview": rows,
            }
        )
    return {"path": str(path.name), "kind": "xlsx/xlsm", "sheets": sheets}


def inspect_xls(path: Path) -> dict:
    import xlrd

    wb = xlrd.open_workbook(path)
    sheets = []
    for sh in wb.sheets():
        rows = []
        for r in range(min(sh.nrows, 12)):
            rows.append([cell_value(sh.cell_value(r, c)) for c in range(min(sh.ncols, 12))])
        sheets.append(
            {
                "title": sh.name,
                "max_row": sh.nrows,
                "max_column": sh.ncols,
                "preview": rows,
            }
        )
    return {"path": str(path.name), "kind": "xls", "sheets": sheets}


def main() -> None:
    results = []
    for path in sorted(ROOT.glob("*.*")):
        if path.suffix.lower() not in {".xls", ".xlsx", ".xlsm"}:
            continue
        try:
            if path.suffix.lower() == ".xls":
                results.append(inspect_xls(path))
            else:
                results.append(inspect_xlsx(path))
        except Exception as exc:
            results.append({"path": str(path.name), "error": repr(exc)})

    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
