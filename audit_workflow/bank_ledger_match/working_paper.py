from __future__ import annotations

import shutil
from copy import copy as copy_style
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .config import output_dir, resolve_path
from .utils import parse_amount, parse_date, text


def fill_working_paper(config: dict[str, Any]) -> Path:
    wp_config = config.get("working_paper", {})
    if not wp_config.get("enabled", True):
        raise RuntimeError("working_paper.enabled=false，跳过底稿填报。")

    template = resolve_path(config, config.get("project", {}).get("template_path"))
    if template is None or not template.exists():
        raise FileNotFoundError(f"找不到底稿模板：{template}")

    output_file = resolve_path(config, wp_config.get("output_file"))
    assert output_file is not None
    output_file.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output_file)

    wb = load_workbook(output_file, keep_vba=output_file.suffix.lower() == ".xlsm")
    if wp_config.get("wp03", {}).get("enabled", True):
        _fill_wp03(wb, config)
    if wp_config.get("wp02", {}).get("enabled", True):
        _fill_wp02(wb, config)
    wb.save(output_file)
    return output_file


def _fill_wp03(wb, config: dict[str, Any]) -> None:
    wp = config.get("working_paper", {}).get("wp03", {})
    matches_path = output_dir(config) / "matches" / "matches.csv"
    matches = pd.read_csv(matches_path, dtype=str).fillna("") if matches_path.exists() else pd.DataFrame()
    if matches.empty:
        return

    sheet = wp.get("sheet", "WP-03")
    ws = wb[sheet]
    start_row = int(wp.get("start_row", 9))
    data_end_row = int(wp.get("data_end_row", 38))
    template_row = int(wp.get("template_row", start_row))
    min_amount = parse_amount(wp.get("min_amount", 0))
    columns = wp.get("columns", {})

    rows = matches[matches["amount"].apply(parse_amount) >= min_amount].copy()
    rows = rows.sort_values(["flow", "amount"], ascending=[True, False])
    required = len(rows)
    capacity = data_end_row - start_row + 1
    if required > capacity:
        ws.insert_rows(data_end_row + 1, required - capacity)
        for row_no in range(data_end_row + 1, data_end_row + 1 + required - capacity):
            _copy_row_style(ws, template_row, row_no)

    for row_no in range(start_row, max(data_end_row, start_row + required - 1) + 1):
        _clear_row(ws, row_no, columns.values())

    project = config.get("project", {})
    entity = project.get("client_short_name") or project.get("client_name", "")
    if wp.get("account_cell") and not rows.empty:
        first = rows.iloc[0]
        ws[wp["account_cell"]] = f"{text(first.get('bank_name'))}  {text(first.get('account_no'))}".strip()

    for offset, (_, match) in enumerate(rows.iterrows()):
        row_no = start_row + offset
        _copy_row_style(ws, template_row, row_no)
        _set(ws, columns.get("entity"), row_no, entity)
        _set(ws, columns.get("bank_date"), row_no, _date_or_text(match.get("bank_date")))
        _set(ws, columns.get("bank_debit"), row_no, _blank_if_zero(match.get("bank_debit"), zero_value=0))
        _set(ws, columns.get("bank_credit"), row_no, _blank_if_zero(match.get("bank_credit"), zero_value=0))
        _set(ws, columns.get("counterparty"), row_no, text(match.get("counterparty_name")))
        _set(ws, columns.get("ledger_date"), row_no, _date_or_text(match.get("ledger_date")))
        _set(ws, columns.get("voucher_no"), row_no, text(match.get("voucher_no")))
        _set(ws, columns.get("ledger_summary"), row_no, text(match.get("ledger_summary")))
        _set(ws, columns.get("ledger_counterparty"), row_no, text(match.get("counterparty_name")))
        _set(ws, columns.get("ledger_debit"), row_no, _blank_if_zero(match.get("ledger_debit")))
        _set(ws, columns.get("ledger_credit"), row_no, _blank_if_zero(match.get("ledger_credit")))


def _fill_wp02(wb, config: dict[str, Any]) -> None:
    wp = config.get("working_paper", {}).get("wp02", {})
    clean_dir = output_dir(config) / "clean"
    bank_path = clean_dir / "bank_transactions.csv"
    ledger_path = clean_dir / "ledger_entries.csv"
    if not bank_path.exists() or not ledger_path.exists():
        return

    bank = pd.read_csv(bank_path, dtype=str).fillna("")
    ledger = pd.read_csv(ledger_path, dtype=str).fillna("")
    sheet = wp.get("sheet", "WP-02")
    ws = wb[sheet]
    start_row = int(wp.get("month_start_row", 7))
    bank_name = _first_value(bank, "bank_name") or _first_value(ledger, "bank_name")
    account_no = _first_value(bank, "account_no") or _first_value(ledger, "account_no")

    for month in range(1, 13):
        row_no = start_row + month - 1
        _set(ws, wp.get("bank_name_col"), row_no, bank_name)
        _set(ws, wp.get("account_no_col"), row_no, account_no)
        _set(ws, wp.get("ledger_debit_col"), row_no, _monthly_sum(ledger, "ledger_debit", month))
        _set(ws, wp.get("bank_credit_col"), row_no, _monthly_sum(bank, "bank_credit", month))
        _set(ws, wp.get("ledger_credit_col"), row_no, _monthly_sum(ledger, "ledger_credit", month))
        _set(ws, wp.get("bank_debit_col"), row_no, _monthly_sum(bank, "bank_debit", month))


def _set(ws, col: str | None, row_no: int, value: Any) -> None:
    if not col:
        return
    col_idx = column_index_from_string(col)
    cell = ws.cell(row=row_no, column=col_idx)
    for merged_range in ws.merged_cells.ranges:
        if cell.coordinate in merged_range:
            cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            break
    cell.value = value


def _clear_row(ws, row_no: int, cols) -> None:
    for col in cols:
        if col:
            ws.cell(row=row_no, column=column_index_from_string(col)).value = None


def _copy_row_style(ws, src_row: int, dst_row: int) -> None:
    if src_row == dst_row:
        return
    for col in range(1, ws.max_column + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy_style(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy_style(src.alignment)
        if src.font:
            dst.font = copy_style(src.font)
        if src.fill:
            dst.fill = copy_style(src.fill)
        if src.border:
            dst.border = copy_style(src.border)


def _date_or_text(value: object) -> object:
    raw = text(value)
    if "|" in raw:
        return raw
    return parse_date(raw) or raw


def _blank_if_zero(value: object, zero_value: object | None = None) -> object:
    amount = parse_amount(value)
    if amount == 0:
        return zero_value
    return amount


def _monthly_sum(df: pd.DataFrame, col: str, month: int) -> float:
    if df.empty or col not in df or "transaction_date" not in df:
        return 0.0
    total = 0.0
    for _, row in df.iterrows():
        trans_date = parse_date(row.get("transaction_date"))
        if trans_date and trans_date.month == month:
            total += parse_amount(row.get(col))
    return round(total, 2)


def _first_value(df: pd.DataFrame, col: str) -> str:
    if df.empty or col not in df:
        return ""
    for value in df[col]:
        item = text(value)
        if item:
            return item
    return ""
